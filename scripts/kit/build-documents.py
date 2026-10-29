"""Build original, editable Guildframe templates and the two reader PDFs.

Run with the bundled Python runtime. QA outputs stay outside the repository.
"""
from pathlib import Path
from xml.sax.saxutils import escape
import json
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from PIL import Image
from pypdf import PdfReader

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'public/downloads/launch-kit'
QA=Path('/private/tmp/guildframe-kit-qa')
OUT.mkdir(parents=True,exist_ok=True); QA.mkdir(parents=True,exist_ok=True)
DATE='29 October 2026'
TEAL='165B55'; DARK='18282D'; GRAY='D9D9D9'

def new_doc(title, opening):
    d=Document(); sec=d.sections[0]
    sec.top_margin=Inches(.67);sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.75);sec.right_margin=Inches(.75)
    normal=d.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(10.5)
    normal.paragraph_format.space_after=Pt(6)
    normal.paragraph_format.line_spacing=1.10
    for name,size in [('Title',25),('Subtitle',11),('Heading 1',17),('Heading 2',12),('Heading 3',11)]:
        st=d.styles[name];st.font.name='Arial';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
        st.paragraph_format.space_before=Pt(10);st.paragraph_format.space_after=Pt(6)
    d.core_properties.author='Guildframe';d.core_properties.title=title;d.core_properties.subject='Editable tabletop crowdfunding templates'
    # The runtime starter can carry a decorative title border. Keep titles plain.
    for node in list(d.styles.element.iter(qn('w:pBdr'))):
        node.getparent().remove(node)
    d.add_paragraph(title,'Title');d.add_paragraph('GUILDFRAME  |  EDITABLE TEMPLATE  |  '+DATE,'Subtitle')
    d.add_paragraph(opening)
    footer=sec.footer.paragraphs[0];footer.alignment=2
    r=footer.add_run('Guildframe  ');r.font.size=Pt(8)
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
    return d

def p(d,text,bold=False):
    para=d.add_paragraph();run=para.add_run(text);run.bold=bold;return para
def h(d,text):d.add_paragraph(text,'Heading 1')
def h2(d,text):d.add_paragraph(text,'Heading 2')
def page(d,title):d.add_page_break();h(d,title)
def bullet(d,text):d.add_paragraph(text,'List Bullet')
def table(d,headers,rows,widths=None):
    t=d.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for i,v in enumerate(headers):t.rows[0].cells[i].text=v
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row):cells[i].text=str(v)
    for ri,row in enumerate(t.rows):
        trPr=row._tr.get_or_add_trPr()
        if ri==0:
            repeat=OxmlElement('w:tblHeader');trPr.append(repeat)
        cant=OxmlElement('w:cantSplit');trPr.append(cant)
        for ci,c in enumerate(row.cells):
            if widths:c.width=Inches(widths[ci])
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),GRAY);b.append(e)
            pr.append(b);mar=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');mar.append(e)
            pr.append(mar)
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),TEAL if ri==0 else ('F1F5F2' if ri%2==0 else 'FFFFFF'));pr.append(sh)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(2);para.paragraph_format.space_before=Pt(2)
                for run in para.runs:
                    run.font.size=Pt(9)
                    if ri==0:run.font.color.rgb=RGBColor(255,255,255);run.bold=True
    d.add_paragraph()
    return t

emails=[
('01 Deliver the kit and welcome the reader','On signup after verified delivery','New subscriber who asked for your updates','Your [PROJECT] preview is here','A first look at what we are making.',
 'Hi [FIRST NAME],\nThanks for joining [PROJECT]. Here is [THE PROMISED DOWNLOAD OR PREVIEW]: [LINK].\nWe are making [SPECIFIC PRODUCT] for [AUDIENCE]. The current version is [HONEST PROTOTYPE STATUS]. Over the next [PERIOD], we will share [WHAT THEY OPTED IN FOR].\nStart with [ONE USEFUL ACTION]. If you have a question about [SPECIFIC TOPIC], reply to this email.\n[CREATOR NAME]',
 'Open the promised resource','Use only the resource promised on your signup page. Keep required delivery separate from optional marketing.'),
('02 Introduce the experience','After the welcome email','Opted in product subscribers','What you do in [PROJECT]','One concrete moment from the experience.',
 'Hi [FIRST NAME],\nIn [PROJECT], you [CONCRETE ACTION]. The interesting choice is [DECISION], because [CONSEQUENCE].\nWe made it for [AUDIENCE OR OCCASION]. A session takes [VERIFIED TIME], with [PLAYER COUNT OR RELEVANT COMPATIBILITY].\nThis [PHOTO, CLIP OR RULES EXCERPT] shows the current version: [LINK]. It is [PROTOTYPE OR FINAL STATUS].\nWhat would you want to see next? Reply with your question.\n[CREATOR NAME]',
 'See the example','Replace game fields with use and fit details for accessories. Do not invent performance claims.'),
('03 Show evidence','After you have something real to demonstrate','Relevant subscribers','See [ONE REAL FEATURE] in use','A clear look at the current prototype.',
 'Hi [FIRST NAME],\nHere is [PROJECT] doing [SPECIFIC THING]: [LINK].\nThe [CLIP, SAMPLE OR TEST] shows [OBSERVABLE DETAIL]. Since the last version, we changed [REAL CHANGE] because [ACTUAL FEEDBACK OR TEST].\nStill in progress: [UNRESOLVED WORK]. We will show [NEXT PROOF] when it is ready.\nYou can inspect [RULES, QUICKSTART, TEST PRINT OR DIMENSIONS] here: [LINK].\n[CREATOR NAME]',
 'Inspect the product proof','Use real tests and permissions. If you do not have evidence, make the evidence before using this draft.'),
('04 Ask a useful question','When a decision can still change','Subscribers familiar with the product','Help us choose [SPECIFIC DECISION]','One question that will change the next version.',
 'Hi [FIRST NAME],\nWe are deciding between [OPTION A] and [OPTION B] for [PROJECT].\n[OPTION A] means [HONEST TRADEOFF]. [OPTION B] means [HONEST TRADEOFF].\nWhich would you choose for [SPECIFIC USE], and why? Reply with A or B and a sentence about how you would use it.\nWe will use the replies alongside [COST, PLAYTEST OR FIT EVIDENCE]. This is feedback, not a promise that either option will be produced.\n[CREATOR NAME]',
 'Reply with a choice and reason','Ask one question. Do not treat a poll response as a purchase commitment.'),
('05 Announce the launch date','After the campaign is approved and date is confirmed','Opted in launch subscribers','[PROJECT] launches [DAY AND DATE]','Launch time and the simplest way to follow.',
 'Hi [FIRST NAME],\n[PROJECT] is scheduled to launch on Kickstarter on [DATE] at [TIME AND TIMEZONE].\nThe campaign will offer [CONFIRMED CORE REWARD] at [CONFIRMED PRICE AND CURRENCY]. [CLEAR SHIPPING AND TAX EXPECTATION].\nYou can follow the project here: [KICKSTARTER PRELAUNCH LINK]. Kickstarter can notify you when it goes live.\nWe will send the live link after we have checked the campaign.\n[CREATOR NAME]',
 'Follow the prelaunch page','Confirm approval, exact timezone, reward price and shipping wording before sending.'),
('06 Send the day before reminder','24 hours before confirmed launch','Launch subscribers who have not opted out','Tomorrow at [TIME]  [PROJECT]','The reward and product details in one place.',
 'Hi [FIRST NAME],\nWe plan to launch [PROJECT] tomorrow, [DATE], at [TIME AND TIMEZONE].\nIf you want to review it first, here are [RULES, PRODUCT PROOF OR QUICKSTART]: [LINK]. The core reward includes [EXACT CONTENTS].\n[SHIPPING AND TAX EXPECTATION]. Any [GENUINE LIMITED OFFER] is described on the campaign; check its actual quantity and terms there.\nWe will send the checked live link tomorrow.\n[CREATOR NAME]',
 'Review the product or follow the page','Remove limited offer wording unless the offer and its limits exist.'),
('07 Announce the live campaign','After live URL, rewards and shipping have been checked','Opted in launch subscribers','[PROJECT] is live on Kickstarter','See the campaign and choose a reward.',
 'Hi [FIRST NAME],\n[PROJECT] is now live: [UNIQUE KICKSTARTER LINK].\nIt is [ONE SENTENCE PRODUCT DESCRIPTION]. The [CORE REWARD] includes [CONTENTS] for [PRICE AND CURRENCY]. [SHIPPING AND TAX EXPECTATION].\nThe page includes [REAL GAMEPLAY OR USE PROOF], the production plan, and the risks we are working through.\nIf the project fits what you enjoy, you can choose a reward on the campaign page. If you have a question, reply here.\n[CREATOR NAME]',
 'View the live campaign','Suppress known backers from repeated acquisition asks when status is available. Check the live page before sending.'),
('08 Share actual progress','When there is a material verified update','Backers or subscribers, with wording adjusted','[PROJECT] update  [SPECIFIC CHANGE]','What changed and what happens next.',
 'Hi [FIRST NAME],\nSince launch, [VERIFIED MILESTONE OR REAL PRODUCT UPDATE]. Thank you to everyone who has [ACTUAL ACTION].\nWe have answered [RECURRING QUESTION] here: [LINK OR SHORT ANSWER].\nThe next step is [SPECIFIC TASK], owned by [PERSON OR TEAM], with our next update planned for [DATE]. [KNOWN UNCERTAINTY IF MATERIAL].\n[BACKER VERSION: Thank you for backing the project.] [SUBSCRIBER VERSION: The campaign details are here: LINK.]\n[CREATOR NAME]',
 'Read the update or campaign','Choose one audience version. Never fabricate funding totals or imply collected cash while only pledges exist.'),
('09 Answer the main question','Midcampaign when a real question recurs','Interested subscribers and backers','A question about [PROJECT]','A direct answer before you decide.',
 'Hi [FIRST NAME],\nA question we have heard is: [REAL QUESTION].\nThe answer is [SPECIFIC FACTUAL ANSWER]. [EXPLAIN LIMIT OR TRADEOFF]. You can see the supporting [DEMONSTRATION, DIMENSION OR RULE] here: [LINK].\nWe have added this to the FAQ so everyone can find it. If your situation differs, reply with [DETAIL NEEDED TO HELP].\nCampaign details: [LINK].\n[CREATOR NAME]',
 'Read the answer','Do not promise future compatibility, shipping destinations or features that have not been costed and confirmed.'),
('10 Send the final two day reminder','48 hours before the actual platform deadline','Relevant nonbacker subscribers','[PROJECT] ends [DAY] at [TIME]','A final chance to review the campaign.',
 'Hi [FIRST NAME],\nThe [PROJECT] Kickstarter campaign is scheduled to end on [DATE] at [TIME AND TIMEZONE].\nIf you are still deciding, start with [BEST REAL DEMONSTRATION] and the [REWARD COMPARISON]: [LINK].\n[SHIPPING AND TAX EXPECTATION]. The current estimated delivery is [DATE OR WINDOW], with the risks explained on the page.\nYou can review the rewards here: [UNIQUE CAMPAIGN LINK].\n[CREATOR NAME]',
 'Review rewards','Use the exact live deadline. Suppress known backers and people who opted out.'),
('11 Send the final day reminder','On the final campaign day','Relevant nonbacker subscribers','[PROJECT] closes today at [TIME]','The deadline and the details you need.',
 'Hi [FIRST NAME],\n[PROJECT] closes today, [DATE], at [TIME AND TIMEZONE].\nThe campaign page has the current reward availability, [PRODUCT PROOF], shipping expectations and delivery plan: [LINK].\n[OPTIONAL FACTUAL SENTENCE ABOUT A REAL CHANGE SINCE THE LAST EMAIL].\nIf you would like to take part, please review the campaign before that deadline. Thank you for following the project and for the feedback you have shared.\n[CREATOR NAME]',
 'View the current campaign','Do not use a false countdown or an expired reward price. Check duplicate reminders across automations.'),
('12 Explain funded next steps','After the funded outcome is verified','Project backers through approved project channels','[PROJECT] funded  next steps','What happens with collection and production.',
 'Hi [FIRST NAME],\n[PROJECT] has finished its Kickstarter campaign with a funded outcome. Thank you for helping make it possible.\nWe are now [ACTUAL COLLECTION OR PREPRODUCTION STEP]. The next tasks are [TASK ONE] and [TASK TWO].\n[IF APPLICABLE: Shipping and applicable taxes will be collected later through PROVIDER. We will explain the amount and timing before collection.]\nOur next project update is planned for [DATE]. Please use [SUPPORT CHANNEL] for reward questions.\n[CREATOR NAME]',
 'Read the next steps','Funding is not delivery. Do not claim final collection, a factory order or a shipment until each is verified.'),
]

d=new_doc('Campaign and Email Templates','Use these editable drafts to turn a product plan into clear campaign copy. Replace bracketed fields with verified project details. The examples describe fictional projects. Keep only the sections that apply to your product and audience.')
h(d,'Your one page campaign plan')
table(d,['Decision','Your project input','Fictional example'],[
 ('Product and audience','[WHAT YOU ARE MAKING AND FOR WHOM]','Harbor Council for players who enjoy negotiation'),
 ('Experience','[ONE SPECIFIC MOMENT OR USE]','Choose which trade agreement to honor this round'),
 ('Current proof','[PROTOTYPE AND LINK]','Current prototype and rules excerpt'),
 ('Core reward','[CONTENTS AND PRICE]','Base game, illustrative $35 USD'),
 ('Shipping and taxes','[WHEN CHARGED AND WHAT IS KNOWN]','Shipping later; confirm zones and tax plan'),
 ('Production budget','[QUOTE, MOQ AND CASH GAP]','Replace fictional workbook inputs first'),
 ('Launch date','[DATE TIME AND TIMEZONE]','Choose after readiness checks'),
 ('Next three actions','[ACTION OWNER DATE]','Obtain quote, test rules, cost parcel'),
],[1.4,2.7,2.65])
p(d,'Messaging sentence: [PROJECT] is a [PRODUCT] for [AUDIENCE] who want [SPECIFIC EXPERIENCE]. You [CONCRETE ACTION], with [OBSERVABLE DIFFERENCE].',True)
p(d,'Example: Harbor Council is a strategy game for players who enjoy negotiating under pressure. You build trade agreements, then decide which promises you can afford to keep. This describes a fictional concept, not an existing product or proven claim.')
page(d,'Campaign page outline')
outline=[('Opening','Name the product and the experience in two sentences. Show a real prototype or finished sample. State the current development status.'),('Who it is for','Give player count, play time and age guidance for games. Give system and adventure scope for TTRPGs. Give scale, license or compatibility for miniatures and accessories.'),('See it in use','Walk through one turn, sample encounter, test print or fit demonstration. Link rules or a quickstart that actually exists.'),('What is included','List exact quantities, materials, files and versions. Label optional add ons separately.'),('Rewards','Show a small comparison. Explain which reward suits which person and what is shared across tiers.'),('Why crowdfunding','Explain what the campaign pays for and what has already been completed. Use the budget you can deliver at the public goal.'),('Production plan','Name the remaining milestones, dependencies, buffers and approval steps. Do not turn an estimate into a guarantee.'),('Shipping and taxes','State supported regions, when shipping is charged, estimated amounts or ranges, and how tax and duties will be handled.'),('Creator and collaborators','Identify the team and relevant work truthfully. Use names, logos and testimonials only with rights and permission.'),('Questions','Answer observed questions about gameplay, files, compatibility, availability, delivery and support.'),('Risks and challenges','Draft the required platform field separately. Describe real production, licensing, cash, logistics and schedule uncertainty plus the current response.'),('Next action','Give a clear invitation to inspect the page or choose a reward. Include a deadline only when confirmed.')]
for title,text in outline:h2(d,title);p(d,text)
page(d,'Rewards and launch readiness')
p(d,'Complete this comparison from the Rewards tab. Keep the combined contents and shipping treatment consistent across the page, video and reward editor.')
table(d,['Reward','Who it suits','Included','Price and later charges'],[('[CORE]','[MAIN AUDIENCE]','[EXACT ITEMS AND QUANTITIES]','[PRICE CURRENCY SHIPPING TAX]'),('[BUNDLE]','[WHY SOMEONE NEEDS TWO OR MORE]','[EXACT ITEMS AND QUANTITIES]','[PRICE CURRENCY SHIPPING TAX]'),('[DIGITAL]','[DIGITAL AUDIENCE]','[FILES LICENSE AND ACCESS]','[PRICE CURRENCY TAX]')],[1.3,1.6,2,1.85])
h2(d,'Reward description draft')
p(d,'[REWARD NAME] includes [EXACT CONTENTS]. Choose this if [RELEVANT USE]. [PHYSICAL: Shipping is charged WHEN, with current estimates HERE.] [DIGITAL: Files will be delivered through METHOD, with LICENSE.] Estimated delivery is [WINDOW]. See the campaign for current risks and any applicable later charges.')
h2(d,'Complete before launch')
for x in ['The campaign goal funds the promised work at the minimum viable production order.','Funding and Cash are checked with real costs, reserves and payment timing.','Every reward has exact contents, cost, quantity limit and delivery estimate.','Current prototype, rules, files or fit evidence are accessible and accurately labelled.','The chosen shipping regions and later charges are clearly explained.','Identity, bank details and required platform review are complete.','Links and email delivery have been checked on phone and desktop.','Email audience, consent and duplicate suppression are checked.','A named person will handle questions and production updates.']:bullet(d,'[ ] '+x)
p(d,'An unfinished item is an action to resolve. A calendar date alone does not make the campaign ready.')
for i in range(0,len(emails),2):
    page(d,'Email drafts '+str(i+1)+' and '+str(i+2))
    for title,timing,audience,subject,preheader,body,cta,guard in emails[i:i+2]:
        h2(d,title);p(d,'When: '+timing+'. Audience: '+audience+'.')
        p(d,'Subject: '+subject,True);p(d,'Preheader: '+preheader)
        for para in body.split('\n'):p(d,para)
        p(d,'Primary action: '+cta+'. Check: '+guard)
page(d,'Outcome and delay updates')
h2(d,'If the campaign did not fund')
p(d,'Subject: Thank you for following [PROJECT]')
p(d,'Hi [FIRST NAME],\nThe [PROJECT] campaign ended without reaching its funding goal. Under Kickstarter funding rules, campaign pledges are not collected when the goal is not met.\nThank you for the time, feedback and support you gave the project. We are reviewing [SPECIFIC FEEDBACK OR COST ISSUE] before deciding on the next step.\n[IF TRUE: If you opted in to future updates, we will share a decision by DATE.] We will not announce a new launch date until [REAL PREREQUISITE].\n[CREATOR NAME]')
h2(d,'If a milestone is delayed')
p(d,'Subject: [PROJECT] update on [DELAYED MILESTONE]')
p(d,'Hi [FIRST NAME],\n[EXACT MILESTONE] is delayed because [FACTUAL CAUSE]. The last completed step is [VERIFIED STATE].\nWe are taking [ACTION], owned by [PERSON]. We expect the next information from [SOURCE] on [DATE]. [IF KNOWN: The revised estimate is WINDOW.] [IF UNKNOWN: We cannot give a reliable new delivery window yet.]\nThis affects [REWARDS OR DESTINATIONS]. Our next update will be [DATE], even if the issue is still open. Reward questions can go to [CHANNEL].\n[CREATOR NAME]')
p(d,'Use a real update commitment you can keep. Explain the impact and next evidence, without implying that a production estimate is a dispatch confirmation.')
page(d,'Reviewer and partner outreach')
h2(d,'Relevant reviewer introduction')
p(d,'Subject: [PROJECT] prototype for your [SPECIFIC SERIES OR AUDIENCE]')
p(d,'Hi [NAME],\nI saw your [SPECIFIC RECENT PIECE] and thought [PROJECT] might fit because [CONCRETE RELEVANCE].\nIt is [ONE SENTENCE DESCRIPTION], currently at [PROTOTYPE STATUS]. The [RULES, QUICKSTART, PHOTOS OR TEST FILE] are here: [LINK]. We are planning [LAUNCH WINDOW, IF CONFIRMED].\nDo you currently consider [PHYSICAL PROTOTYPES OR DIGITAL SAMPLES] in this category? If so, please share your submission policy, availability and any costs or disclosure requirements.\nThere is no expectation of positive coverage. We would agree the scope and any timing before sending a sample.\n[NAME AND BUSINESS CONTACT]')
h2(d,'One polite follow up')
p(d,'Hi [NAME],\nFollowing up once on [PROJECT], in case it fits your current schedule. The short overview is here: [LINK]. If this is not a fit or the timing is wrong, no problem. We will leave it here.\n[NAME]')
h2(d,'Partner announcement draft')
p(d,'[PROJECT] is [VERIFIED CURRENT STATUS]. It offers [CONCRETE DESCRIPTION] for [AUDIENCE]. You can inspect [PRODUCT PROOF] and the current reward details at [UNIQUE LINK]. [DISCLOSE PAID OR AFFILIATE RELATIONSHIP WHERE APPLICABLE].')
p(d,'Keep contact provenance, policy, cost, prototype shipping and permissions in Reviewers. Confirm any embargo explicitly. Do not reuse a quote or logo without permission, and do not treat a paid placement as an independent endorsement.')
d.save(OUT/'Campaign-and-Email-Templates.docx')

d=new_doc('Creative and Production Briefs','Give collaborators a clear product story, source assets and acceptance criteria. These briefs are editable working requests. Replace bracketed fields, delete inapplicable items, and agree dates and scope before work starts.')
h(d,'Project and asset brief')
table(d,['Field','Your brief'],[('Project and owner','[PROJECT NAME CREATOR CONTACT AND APPROVER]'),('Product and stage','[WHAT EXISTS TODAY AND WHAT REMAINS]'),('Audience and action','[WHO THIS IS FOR AND THE ONE ACTION WE WANT]'),('Main message','[ONE SPECIFIC PROMISE THAT THE PRODUCT CAN SUPPORT]'),('Evidence','[REAL DEMO RULES TEST SAMPLE OR DIMENSIONS]'),('Required assets','[PHOTO VIDEO GRAPHIC OR COPY OUTPUTS]'),('Source files','[LINK TO CURRENT FILES WITH VERSIONS AND RIGHTS]'),('Formats','[EDITABLE ORIGINAL EXPORT SIZE AND PLATFORM]'),('Deadline','[DATE TIMEZONE DEPENDENCIES AND REVIEW ROUNDS]'),('Acceptance','[WHAT MUST BE VISIBLE ACCURATE AND EDITABLE]')],[1.45,5.3])
h2(d,'Worked message example')
p(d,'Fictional example: Harbor Council is for strategy players who enjoy negotiated decisions. Show one turn where a player chooses which trade promise to keep. Use current prototype footage and an overhead table view. The primary action is to read the rules excerpt. Do not claim that the game has shipped, won an award or reached a funding result.')
page(d,'Asset checklist and video storyboard')
table(d,['Asset','Required content','Acceptance check'],[('Opening image','Readable product and one clear idea','Recognizable at phone size; prototype labelled'),('Contents view','Every included component or file','Quantities match reward descriptions'),('Use demonstration','One turn, encounter, print or fit test','Actual product behavior is visible'),('Reward comparison','Core, bundle and digital options','Exact contents, price and shipping timing'),('Production timeline','Remaining milestones and dependencies','Estimate language and next update date'),('Shipping panel','Supported zones and later charges','Matches current provider setup')],[1.4,2.7,2.65])
h2(d,'Six frame video storyboard')
table(d,['Frame','Picture','Narration or caption'],[('1','Real product in use','[WHAT YOU CAN DO WITH IT]'),('2','The specific decision or use','[WHY THIS MOMENT IS INTERESTING]'),('3','Close view of components or file','[WHAT IS INCLUDED]'),('4','Uncut short demonstration','[HOW THE CORE ACTION WORKS]'),('5','Creator and current prototype','[WHAT IS READY AND WHAT FUNDING ENABLES]'),('6','Reward and campaign link','[CLEAR NEXT ACTION AND LATER CHARGES]')],[.55,2.4,3.8])
p(d,'Confirm current platform dimensions before exporting. Deliver editable originals and the final exports. List licensed fonts, music, art, model releases and any AI use that requires disclosure.')
page(d,'Four original ad test briefs')
tests=[('Gameplay or use','A player must choose which promise to keep.','Show a real turn and its consequence. For an accessory, show setup and the useful action.','Audience already interested in this kind of experience.','Read the rules or inspect the use demo.'),('Product detail','What is actually inside the reward.','Show component count, scale, finish or file contents. Match the real prototype and reward.','People comparing products in this category.','See contents and current prototype.'),('Creator process','The change we made after testing.','Show the before and current version, with actual feedback and no invented praise.','People interested in design and development.','Read the development note.'),('Specific audience benefit','A reason this fits your table.','Show a genuine use case. Use verified player count, time, scale or compatibility.','One clearly described audience segment.','See whether it fits your group or setup.')]
for title,hook,shot,audience,cta in tests:
    h2(d,title);p(d,'Hook: '+hook);p(d,'Asset: '+shot);p(d,'Audience hypothesis: '+audience);p(d,'Action: '+cta)
    p(d,'Fill before launch: [SPEND CAP] [START AND STOP DATE] [MEANINGFUL QUALIFIED ACTION] [DECISION RULE] [OWNER]. Change one main variable in the next test.')
p(d,'Track spend, clicks, unique opt ins and qualified actions in Ad tests. Compare the same reporting window. A low CPL does not prove that people will back the campaign.')
page(d,'Manufacturing quote request')
p(d,'Subject: Quote request for [PROJECT]  [PRODUCT TYPE AND QUANTITY]')
p(d,'Hi [SUPPLIER],\nWe are preparing [PROJECT], a [PRODUCT TYPE]. Please quote the specification below and identify assumptions, exclusions and optional changes. We need a written quote before setting our campaign budget.\nProduct specification: [COMPONENT BOM, MATERIALS, DIMENSIONS, PAGE OR CARD COUNT, FINISH, COLOR, SCALE].\nQuantities: [MOQ AND PRICE BREAKS TO QUOTE].\nFiles and stage: [CURRENT DRAWINGS, FILES, PROTOTYPE OR SAMPLE].\nDestination: [COUNTRY, POSTCODE, WAREHOUSE OR PORT].\nPlease include sample and proof cost, tooling, unit price, packing, carton dimensions and weights, order increments, production lead time, payment milestones, quality inspection, defect/rework policy, currency, Incoterm, quote expiry and excluded freight, duties or taxes.\nPlease flag any specification that creates avoidable cost or schedule risk and price alternatives separately.\n[NAME AND CONTACT]')
h2(d,'Compare like for like')
for x in ['Is the quoted unit a complete reward, one component or one carton?','Are tooling, proofing, packaging, inbound freight and duties included?','Does the MOQ apply per material, language, finish, color or SKU?','Which sample must be approved before the production balance is due?','What happens if an inspection fails or units arrive damaged?','Which milestone dates are estimates and what must happen first?']:bullet(d,x)
p(d,'Copy the agreed figures into Costs. Record the quote version and date. The workbook only covers one manufactured SKU shared by tiers; additional independent SKUs require additional complete cost models.')
page(d,'Production and fulfillment handoff')
table(d,['Handoff item','Complete before transfer'],[('Product and rewards','[SKU TO REWARD MAP WITH QUANTITIES AND VERSIONS]'),('Approved source','[SIGNED SAMPLE AND FINAL FILE LOCATION]'),('Order and cash','[QUANTITY SPARES DEPOSIT BALANCE AND CASH EVIDENCE]'),('Quality','[COUNTS DIMENSIONS FINISH DEFECT TOLERANCE AND OWNER]'),('Packout','[PARCEL PHOTO WEIGHT DIMENSIONS AND BUNDLE VARIANTS]'),('Shipping and taxes','[REGIONS COLLECTION TIMING PROVIDER AND RESPONSIBILITY]'),('Addresses','[SURVEY DATE LOCK DATE ACCESS OWNER AND REMINDERS]'),('Pilot shipment','[TEST TRACKING CONTENTS LABEL AND DELIVERY CHECK]'),('Exceptions','[LOST DAMAGED DIGITAL ACCESS OR WRONG ITEM RESPONSE]'),('Communication','[SUPPORT CONTACT NEXT UPDATE DATE AND ESCALATION]')],[1.55,5.2])
h2(d,'Category specific additions')
for x in ['Board games: component count, rulebook version, carton weight and full reward packout.','Card games: deck count, card stock, finish, language and print run price break.','TTRPGs: writing/editing/layout status, page count, accessible PDF, print proof and license.','Miniatures: physical versus STL, scale, base, sculpt rights, test print, supports and breakage.','Accessories: material, tolerance, finish, compatibility dimensions and fit/durability test.']:bullet(d,x)
p(d,'Keep backer names, addresses and payment details in the selected fulfillment system. Share only the data each collaborator needs. Do not paste a backer list into this reusable brief.')
d.save(OUT/'Creative-and-Production-Briefs.docx')

# PDF layout uses natural flowing paragraphs and explicit logical page breaks.
styles=getSampleStyleSheet()
styles.add(ParagraphStyle(name='GFTitle',fontName='Helvetica-Bold',fontSize=26,leading=31,textColor=colors.HexColor('#18282D'),spaceAfter=17))
styles.add(ParagraphStyle(name='GFHead',fontName='Helvetica-Bold',fontSize=15,leading=19,spaceBefore=13,spaceAfter=8,textColor=colors.black))
styles.add(ParagraphStyle(name='GFBody',fontName='Helvetica',fontSize=10.5,leading=15,spaceAfter=9,textColor=colors.HexColor('#253832')))
styles.add(ParagraphStyle(name='GFSmall',fontName='Helvetica',fontSize=8.5,leading=12,spaceAfter=7,textColor=colors.HexColor('#53665F')))
styles.add(ParagraphStyle(name='GFCell',fontName='Helvetica',fontSize=9,leading=12,spaceAfter=0))
styles.add(ParagraphStyle(name='GFCellHead',fontName='Helvetica-Bold',fontSize=9,leading=12,textColor=colors.white))
def P(text,style='GFBody'):
    markup=escape(text).replace('\n','<br/>')
    markup=re.sub(r'(https://[^\s<]+)',lambda m:'<link href="'+m.group(1)+'">'+m.group(1)+'</link>',markup)
    return Paragraph(markup,styles[style])
def H(text):return P(text,'GFHead')
def PT(headers,rows,widths):
    data=[[P(x,'GFCellHead') for x in headers]]+[[P(str(x),'GFCell') for x in row] for row in rows]
    t=Table(data,colWidths=widths,repeatRows=1,hAlign='LEFT');t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#165B55')),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#F1F5F2')]),('GRID',(0,0),(-1,-1),.4,colors.HexColor('#D9D9D9')),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('LEFTPADDING',(0,0),(-1,-1),9),('RIGHTPADDING',(0,0),(-1,-1),9),('TOPPADDING',(0,0),(-1,-1),8),('BOTTOMPADDING',(0,0),(-1,-1),8)]));return t
def footer(canvas,doc):
    canvas.saveState();canvas.setFont('Helvetica',8);canvas.setFillColor(colors.HexColor('#53665F'));canvas.drawString(45,28,'GUILDFRAME  |  Tabletop Kickstarter Launch Kit  |  '+DATE);canvas.drawRightString(A4[0]-45,28,str(doc.page));canvas.restoreState()
def pdf(name,story):
    SimpleDocTemplate(str(OUT/name),pagesize=A4,rightMargin=45,leftMargin=45,topMargin=43,bottomMargin=47,title=name.replace('-',' '),author='Guildframe').build(story,onFirstPage=footer,onLaterPages=footer)

story=[P('Tabletop Kickstarter\nLaunch Kit','GFTitle'),P('Start here','GFHead'),P('Price the rewards, check the cash, and turn the remaining work into a dated launch plan. This kit is for board games, card games, TTRPGs, miniatures and accessories.'),P('Start with a copy of the workbook. It contains a complete fictional example so you can see the formulas work. Replace blue inputs with your project details before relying on any result.'),H('Your five working files'),PT(['File','Use it for'],[('Tabletop-Launch-Planner.xlsx','Reward economics, MOQ, shipping, cash timing and launch trackers'),('Campaign-and-Email-Templates.docx','One page plan, campaign outline, rewards, 12 emails and outreach drafts'),('Creative-and-Production-Briefs.docx','Asset brief, storyboard, ad tests, quote request and fulfillment handoff'),('Examples-and-Sources.pdf','Five category examples, a worked financial case and official references'),('START-HERE.pdf','The order to use the kit and important model boundaries')],[218,287]),H('Your first thirty minutes'),P('1. Save a project copy. In Start, enter your project, currency, launch date and active case. Read the assumptions beside each rate.'),P('2. In Costs, replace the unit cost, MOQ, order increment and fixed costs. In Rewards, replace prices, mix, units and other costs.'),P('3. In Shipping, replace zones, mix and parcel quotes. Choose Later, At pledge or Included on Start.'),P('4. Read Funding and Cash. Resolve a cash gap, missing input or unconfirmed quote before choosing a public goal.'),PageBreak(),P('Read the model correctly','GFTitle'),PT(['Tab','Your next decision'],[('Start','Which assumptions and case are active?'),('Funding','Can the whole order and reserves be funded?'),('Rewards','Does each tier contribute after its costs?'),('Costs','What must be paid and what does the quote exclude?'),('Shipping','Which zones are supported and when are charges collected?'),('Cash','Is money available before each payment is due?'),('Launch plan','Who owns each dependency and dated task?'),('Ad tests','What evidence would justify spending again?'),('Sources','What does each reporting source actually attribute?'),('Reviewers','Which outreach is relevant, agreed and permitted?'),('Live ops','What needs action and what has been verified?'),('Production','Which version, sample and handoff are approved?')],[115,390]),H('What this starter workbook covers'),P('One manufactured product SKU can be used by several physical reward tiers, alongside digital rewards with zero physical units. The example is one base game, a two copy bundle and a digital print and play. Additional independent SKUs need complete separate costing.'),P('All amounts use the campaign currency. Convert external quotes explicitly before entering them. The model does not determine tax liability, calculate tariffs, integrate with a supplier, or prove demand.'),P('One weighted parcel cost is used per physical backer. A larger bundle may need a larger parcel. Enter the additional cost in the reward other cost or use confirmed tier specific shipping quotes.'),PageBreak(),P('Funding is also a cash decision','GFTitle'),P('A positive margin can still leave you short of cash. Minimum orders require paying for units that may remain unsold. Pledges are not money available before payout. Shipping collected later cannot pay an earlier factory deposit.'),H('A fictional MOQ example'),P('Suppose a reward is $35, manufacturing is $14 per unit, other reward cost is $3 per backer, fixed work is $6,000 and the supplier MOQ is 500. This example uses US standard 5% platform plus 3% and $0.30 processing, with no shipping, failures, reserves or opening cash.'),PT(['Backers','Net receipts','Costs with MOQ','Surplus or gap'],[('403','$12,855.70','$14,209.00','($1,353.30)'),('450','$14,355.00','$14,350.00','$5.00')],[70,145,145,145]),P('A simple per reward margin suggests 403 backers. The full order model shows the $1,353.30 gap. Even 450 backers leaves only $5 before contingency or other uncertainty. These are arithmetic examples, not a recommended campaign budget.'),H('The checks to complete'),P('Confirm that reward and shipping mixes total 100%. Confirm that prices and unit quantities are nonnegative, order increments are at least one, and every required input is present. Obtain real production and shipping quotes. Set a tax plan, reserve and realistic payment schedule.'),P('In Cash, blue date fields can be replaced with agreed milestone dates. Keep the event rows chronological. Add a committed extra receipt or payment in the extra rows. The minimum balance identifies the funding needed before those events.'),PageBreak(),P('Move from plan to launch','GFTitle'),H('Choose your category path'),PT(['Category','Start with this proof'],[('Board game','Rules, one observed turn, full component BOM and packed weight'),('Card game','Deck specification, print run cost and parcel economics'),('TTRPG','Quickstart, editorial milestones, PDF access and print proof'),('Miniatures','Physical or STL scope, scale, test print, rights and packout'),('Accessories','Dimensions, compatibility, fit test and material specification')],[110,395]),H('Use the copy with verified details'),P('The editable drafts contain real sample wording and bracketed fields. Replace every field before sending. Delete unsupported statements. Keep dates, prices, contents and shipping wording consistent with the live campaign.'),P('Use newsletters only for people who opted in. Keep required reward and delivery communication separate. Use a reviewer business contact and submission policy. Record permission before reusing quotes, logos or footage.'),H('Before you commit money or go live'),P('Complete the launch readiness checklist in the campaign document. Resolve product proof, rights, quotes, cash timing, bank and identity checks, platform review, support ownership and email delivery. A planned date alone does not make a campaign ready.'),H('Where to read more'),P('Guildframe guides: https://guildframe.com/guides'),P('Kickstarter Creator Handbook: https://www.kickstarter.com/help/handbook'),P('Official fee, shipping, tax, attribution and rules references are in Examples and Sources. Check current country and provider terms before setting final figures.'),P('Version 1.0. Free kit. Guildframe is an independent provider and is not affiliated with Kickstarter. The fictional examples describe planning decisions and do not represent customers or campaign results.','GFSmall')]
pdf('START-HERE.pdf',story)

examples=[
 ('Board game','Harbor Council','A fictional strategy game about negotiating trade promises.','Start with a real rules excerpt and one complete turn. Show the decision and its consequence. The audience is players who like negotiation, not everyone who likes board games.', [('Core reward','One base game with the exact listed components'),('Bundle','Two base games; check the larger parcel quote'),('Digital','Print and play files with a clear personal use license')], 'Build a component BOM, including rulebook, insert, packaging and spare parts. Obtain MOQ and landed price per complete base game. Record a packed sample weight and the shipping quote by destination.', 'Do not announce a stretch component until you calculate its cost for every eligible backer, tooling, freight and packout. More pledges can create more obligations.', 'Show the current prototype and label renders. Record playtest issues, the rules version, and the remaining approval steps.'),
 ('Card game','Pocket Rivals','A fictional compact card game for quick sessions.','Show setup and a short turn. Demonstrate the actual choice in the game and publish the current rules. Verify the claimed player count and play time through testing.', [('Core reward','One deck with stated card count and stock'),('Bundle','Two or more decks with exact quantities'),('Language edition','A separate SKU only if its print run is fully costed')], 'Ask for card stock, finish, deck count, tuck box, seal, carton quantity, proof cost and print run price breaks. For a low priced reward, postage and fixed payment fees can materially change contribution.', 'A translated edition may have its own MOQ, editing cost and inventory. Do not divide one manufacturing quote across language variants that the supplier priced separately.', 'Compare an order of one deck with a bundle, including parcel size. A higher average pledge is useful only when the full reward and shipping still contribute.'),
 ('TTRPG','Ash Lantern','A fictional original adventure about a town that loses its shadows.','Offer an actual quickstart, sample spread or encounter. Explain the system, level range or standalone rules clearly. Record rights to text, art, fonts and licensed material.', [('Digital PDF','Accessible PDF and stated personal use license'),('Print edition','Printed book with dimensions and binding'),('Print and PDF','Both versions with clear delivery methods')], 'Budget writing, editing, development, layout, art and file accessibility as separate milestones. Obtain print proof, page count and binding quotes. Set zero physical units for a digital only reward.', 'A PDF has no physical postage or factory MOQ, but still has editorial and delivery costs. A print reward needs a complete separate print and parcel budget.', 'Test PDF download, links, reading order and device access. Name which sections are finished and who approves the final proof.'),
 ('Miniatures','Stonewatch Patrol','A fictional set of original fantasy figures.','Show a test print at the intended scale, including base dimensions and close details. State whether the reward is physical miniatures or STL files. A render is not a manufactured sample.', [('Physical set','Stated number of figures, material and bases'),('Digital STL','File list, supported or unsupported status and license'),('Larger set','Exact extra figures and the corresponding packout')], 'Separate sculpting and rights from master, mold, casting or print costs. Include support testing, failed prints, breakage, packaging and replacement stock. Obtain quotes for every physical variant.', 'A personal STL license does not imply commercial printing rights. State the actual license. Physical and digital rewards require different production and fulfillment steps.', 'Approve scale, fit, supports, orientation and visible quality before batch production. Test a packed parcel for transit damage.'),
 ('Accessories','Foldline Tray','A fictional modular dice tray for small tables.','Show dimensions, setup and a real fit test. Explain the material and intended use. Claim compatibility only where you have checked the relevant dimensions and conditions.', [('Core tray','One tray with stated material, size and finish'),('Pair','Two trays with exact colors or finishes'),('Optional insert','Separate add on only after tooling and fit are costed')], 'Quote material, tolerance, finish, tooling, sample approval, packaging and order increments. Color and finish choices can multiply SKUs and supplier minimums.', 'A new finish may require a separate minimum order and later delivery. Cost the variant before promising it as an unlocked option.', 'Record tolerance and fit acceptance, surface finish, wear testing and packout. Photograph the current sample and list remaining work.'),
]
story=[P('Examples and Sources','GFTitle'),P('Five fictional projects show how the same planning tools change with the product. These are original teaching examples, not real customers, manufacturer quotes, benchmarks or successful campaigns.'),H('Use an example to make a decision'),P('Choose the closest category, then replace the assumptions with your own product evidence and written quotes. The workbook starts with the fictional Harbor Council example. The other paths explain which costs, proof and delivery steps need different treatment.'),H('The MOQ calculation'),P('For a $35 reward with 5% platform and 3% plus $0.30 processing, net campaign receipts are $31.90 per collected pledge. Subtracting $14 manufacturing and $3 other cost gives a simple $14.90 contribution. Dividing $6,000 fixed work by that margin suggests 403 backers.'),P('The factory requires 500 units. At 403 backers, net receipts are $12,855.70, while $6,000 fixed work, $7,000 manufacturing and $1,209 other costs total $14,209. The gap is $1,353.30. At 450 backers, net receipts of $14,355 barely exceed $14,350 costs. Fifty units remain in stock, not in cash.'),P('This illustration excludes shipping, failed payments, taxes, refunds, marketing uncertainty and contingency. It shows why the real goal needs a complete delivery budget and a cash calendar.'),H('Rules for the examples'),P('An estimate is an estimate. A quote needs scope, date, currency and expiry. A positive contribution does not fund the MOQ automatically. Later collections cannot finance earlier payments. Proof should show the actual product state. No page, email or example should promise an outcome you cannot substantiate.')]
for cat,name,desc,proof,rewards,costs,trap,action in examples:
    story += [PageBreak(),P(cat+' example','GFTitle'),H(name),P(desc+' All details on this page are hypothetical.'),H('Evidence before the pitch'),P(proof),PT(['Reward','What must be clear'],rewards,[115,390]),H('Cost and delivery decisions'),P(costs),H('A mistake to catch'),P(trap),H('Next practical action'),P(action)]
story += [PageBreak(),P('First party sources','GFTitle'),P('Checked '+DATE+'. Platform rules and provider terms can change. Open the current country specific page before committing to a price or a collection method.'),H('Kickstarter funding and fees'),P('Funding goal and budgeting: Kickstarter says the goal should cover the minimum needed to make the project and fulfill rewards. The campaign uses all or nothing funding.'),P('https://www.kickstarter.com/help/handbook/funding','GFSmall'),P('US fee defaults: 5% platform fee, with 3% plus $0.30 standard processing per pledge. Pledges under $10 use 5% plus $0.08 processing. Other countries differ.'),P('https://www.kickstarter.com/help/fees?country=US','GFSmall'),H('Shipping and taxes'),P('Kickstarter Pledge Manager permits shipping at pledge time or collection later. Tax handling depends on the jurisdiction and setup. The workbook later fee inputs are illustrative and must be replaced with current provider terms.'),P('https://start.kickstarter.com/pledge-manager-release-notes','GFSmall'),P('https://updates.kickstarter.com/navigating-tax-collection-with-kickstarters-pledge-manager/','GFSmall'),H('Promotion and responsibilities'),P('Kickstarter offers custom referral tags before launch through Promotions. Use platform reported attribution without claiming that every delayed or cross device pledge can be identified.'),P('https://updates.kickstarter.com/new-tools-to-help-creators-market-their-projects/','GFSmall'),P('Present the product honestly, respect rights and backer privacy, explain AI use where required, and avoid unsolicited messages. Fulfillment contact details do not automatically become newsletter permission.'),P('https://www.kickstarter.com/rules','GFSmall'),P('https://legal.kickstarter.com/','GFSmall'),P('Guildframe guides: https://guildframe.com/guides','GFSmall')]
pdf('Examples-and-Sources.pdf',story)

result={}
for f in OUT.glob('*.pdf'):
    reader=PdfReader(f);result[f.name]={'pages':len(reader.pages),'characters':sum(len(p.extract_text() or '') for p in reader.pages)}
for f in OUT.glob('*.docx'):
    doc=Document(f);text=' '.join(p.text for p in doc.paragraphs);result[f.name]={'paragraphs':len(doc.paragraphs),'tables':len(doc.tables),'characters':len(text)}
(QA/'document-structure.json').write_text(json.dumps(result,indent=2))
print(json.dumps(result))
