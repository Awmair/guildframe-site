"""Validate and package the five downloadable launch kit files.

Run after build-workbook.mjs and build-documents.py with bundled Python.
Only the five working files enter the ZIP. The public manifest records hashes.
"""
from pathlib import Path
import hashlib
import json
import zipfile
import xml.etree.ElementTree as ET
from pypdf import PdfReader
from docx import Document

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'public/downloads/launch-kit'
QA = Path('/private/tmp/guildframe-kit-qa')
NAMES = [
    'START-HERE.pdf',
    'Tabletop-Launch-Planner.xlsx',
    'Campaign-and-Email-Templates.docx',
    'Creative-and-Production-Briefs.docx',
    'Examples-and-Sources.pdf',
]
NS = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
manifest = {'version': '1.0', 'created': '2026-10-29', 'files': []}
checks = {}

for name in NAMES:
    path = OUT / name
    data = path.read_bytes()
    assert len(data) > 5000, f'Empty or incomplete file: {name}'
    manifest['files'].append({'name': name, 'bytes': len(data), 'sha256': hashlib.sha256(data).hexdigest()})
    if path.suffix == '.pdf':
        reader = PdfReader(path)
        count = 4 if name == 'START-HERE.pdf' else 7
        assert len(reader.pages) == count
        assert all(len(page.extract_text() or '') > 200 for page in reader.pages)
        checks[name] = {'pages': count, 'textOnEveryPage': True}
    else:
        with zipfile.ZipFile(path) as z:
            assert z.testzip() is None
            assert not any('vbaProject' in p for p in z.namelist())
            if path.suffix == '.xlsx':
                wb = ET.fromstring(z.read('xl/workbook.xml'))
                names = [s.attrib['name'] for s in wb.findall('s:sheets/s:sheet', NS)]
                assert names == ['Start', 'Funding', 'Rewards', 'Costs', 'Shipping', 'Cash', 'Launch plan', 'Ad tests', 'Sources', 'Reviewers', 'Live ops', 'Production']
                formulas = 0
                for p in z.namelist():
                    if p.startswith('xl/worksheets/sheet') and p.endswith('.xml'):
                        sheet = ET.fromstring(z.read(p))
                        formulas += len(sheet.findall('.//s:f', NS))
                        assert not sheet.findall('.//s:c[@t="e"]', NS), f'Formula error: {p}'
                assert formulas == 189
                checks[name] = {'tabs': names, 'formulas': formulas, 'cachedFormulaErrors': 0}
            else:
                settings = ET.fromstring(z.read('word/settings.xml'))
                assert not settings.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}documentProtection')
                doc = Document(path)
                checks[name] = {'paragraphs': len(doc.paragraphs), 'tables': len(doc.tables), 'editable': True}

bundle = OUT / 'tabletop-kickstarter-launch-kit.zip'
with zipfile.ZipFile(bundle, 'w', zipfile.ZIP_DEFLATED, compresslevel=9) as z:
    for name in NAMES:
        info = zipfile.ZipInfo(name, date_time=(2026, 10, 29, 0, 0, 0))
        info.compress_type = zipfile.ZIP_DEFLATED
        info.external_attr = 0o644 << 16
        z.writestr(info, (OUT / name).read_bytes(), compress_type=zipfile.ZIP_DEFLATED, compresslevel=9)
with zipfile.ZipFile(bundle) as z:
    assert z.namelist() == NAMES
    assert z.testzip() is None
    for entry in manifest['files']:
        assert hashlib.sha256(z.read(entry['name'])).hexdigest() == entry['sha256']
manifest['bundle'] = {'name': bundle.name, 'bytes': bundle.stat().st_size, 'sha256': hashlib.sha256(bundle.read_bytes()).hexdigest()}
(OUT / 'manifest.json').write_text(json.dumps(manifest, indent=2) + '\n')
QA.mkdir(exist_ok=True, parents=True)
(QA / 'package-checks.json').write_text(json.dumps({'checks': checks, 'bundle': manifest['bundle'], 'zipEntries': NAMES}, indent=2) + '\n')
print(json.dumps({'bundle': manifest['bundle'], 'files': len(NAMES), 'checks': 'passed'}))
